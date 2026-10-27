"""Таблицы CSV / TSV / Excel (.xlsx): каждая строка таблицы — отдельный текст.

Пользователь выбирает столбец с текстом и (необязательно) столбец с названием;
по умолчанию они угадываются: текст — столбец с самыми длинными значениями,
название — столбец с заголовком вроде «Название», «Товар», «ID».
"""

from __future__ import annotations

import csv
import io
import re
from collections import Counter
from dataclasses import dataclass, field, replace
from datetime import date, datetime, time
from pathlib import Path

from .loaders import TABLE_EXTENSIONS, UnsupportedFileError, read_text_file

csv.field_size_limit(2**31 - 1)  # длинные тексты в ячейках CSV больше стандартного лимита 128 КБ

CSV_DELIMITERS = ";,\t|"
DELIMITER_NAMES = {";": "точка с запятой", ",": "запятая", "\t": "табуляция", "|": "вертикальная черта"}
# начала слов в заголовке, по которым угадываем столбец с названием текста
TITLE_HINTS = ("назван", "заголов", "тема", "товар", "продукт", "имя", "автор", "респондент",
               "title", "name", "subject", "product", "author", "id", "№", "номер", "код")


@dataclass
class Table:
    path: Path
    rows: list[list[str]]
    first_row: int = 1                       # номер строки файла/листа, с которой начинается rows
    sheet: str | None = None
    sheet_names: list[str] = field(default_factory=list)
    encoding: str | None = None
    delimiter: str | None = None

    @property
    def n_cols(self) -> int:
        return max((len(r) for r in self.rows), default=0)

    @property
    def kind(self) -> str:
        return TABLE_EXTENSIONS.get(self.path.suffix.lower(), "таблица")

    def describe(self) -> str:
        parts = [self.kind]
        if self.sheet and len(self.sheet_names) > 1:
            parts.append(f"лист «{self.sheet}»")
        if self.encoding:
            parts.append(f"кодировка {self.encoding}")
        if self.delimiter:
            parts.append(f"разделитель — {DELIMITER_NAMES.get(self.delimiter, self.delimiter)}")
        elif self.path.suffix.lower() in (".csv", ".tsv"):
            parts.append("один столбец")
        parts.append(f"строк: {len(self.rows)}")
        return ", ".join(parts)


@dataclass
class TableSpec:
    text_col: int
    title_col: int | None = None
    has_header: bool = True
    sheet: str | None = None
    # заголовки выбранных столбцов — чтобы применить тот же выбор к другим таблицам
    text_name: str = ""
    title_name: str = ""


# --- чтение ------------------------------------------------------------------

def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d") if value.time() == time(0) else value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value)


def _trim(rows: list[list[str]]) -> tuple[list[list[str]], int]:
    """Убирает пустые строки в начале и конце и пустые столбцы справа. Возвращает (строки, сдвиг)."""
    def empty(row):
        return not any(c.strip() for c in row)

    start = 0
    while start < len(rows) and empty(rows[start]):
        start += 1
    end = len(rows)
    while end > start and empty(rows[end - 1]):
        end -= 1
    rows = rows[start:end]
    last = max((max((i for i, c in enumerate(r) if c.strip()), default=-1) for r in rows), default=-1)
    return [r[: last + 1] for r in rows], start


def _sniff_delimiter(text: str) -> str | None:
    """Разделитель, при котором почти у всех строк одинаковое число столбцов (≥ 2).

    Запятые и точки с запятой внутри обычного текста дают «рваные» строки, поэтому
    файл, где такого разделителя нет, считается таблицей из одного столбца (None)."""
    lines = text.splitlines()
    sample = "\n".join(lines[:200])
    truncated = len(lines) > 200
    best, best_key = None, (0.0, 0)
    for d in CSV_DELIMITERS:
        rows = list(csv.reader(io.StringIO(sample), delimiter=d))
        if truncated:
            rows = rows[:-1]  # последняя запись могла оборваться посреди многострочной ячейки
        counts = [len(r) for r in rows if any(c.strip() for c in r)]
        if not counts:
            continue
        n_fields, freq = Counter(counts).most_common(1)[0]
        consistency = freq / len(counts)
        if n_fields >= 2 and consistency >= 0.8 and (consistency, n_fields) > best_key:
            best, best_key = d, (consistency, n_fields)
    return best


def read_csv(path: Path) -> Table:
    text, encoding = read_text_file(path)
    text = text.replace("\r\n", "\n").replace("\r", "\n")  # переносы внутри ячеек в кавычках — тоже
    delimiter = "\t" if path.suffix.lower() == ".tsv" else _sniff_delimiter(text)
    if delimiter is None:
        rows = [[line] for line in text.splitlines()]  # один столбец: каждая строка — текст
    else:
        rows = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]
    rows, offset = _trim(rows)
    return Table(path, rows, first_row=offset + 1, encoding=encoding, delimiter=delimiter)


def read_xlsx(path: Path, sheet: str | None = None) -> Table:
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException
    from zipfile import BadZipFile

    try:
        wb = openpyxl.load_workbook(path, data_only=True)  # data_only: значения формул, а не сами формулы
    except (BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        raise UnsupportedFileError(
            "файл повреждён или это не книга Excel .xlsx (старый формат .xls сохраните в Excel как .xlsx)"
        ) from exc
    try:
        names = list(wb.sheetnames)
        name = sheet if sheet in names else wb.active.title
        ws = wb[name]
        rows = [[_cell_to_str(v) for v in r] for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    rows, offset = _trim(rows)
    return Table(path, rows, first_row=offset + 1, sheet=name, sheet_names=names)


def read_table(path: str | Path, sheet: str | None = None) -> Table:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return read_xlsx(path, sheet)
    return read_csv(path)


# --- столбцы -----------------------------------------------------------------

def column_letter(index: int) -> str:
    """0 -> A, 25 -> Z, 26 -> AA — как в Excel."""
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(ord("A") + rem) + letters
    return letters


def header_cell(table: Table, has_header: bool, col: int) -> str:
    if not has_header or not table.rows or col >= len(table.rows[0]):
        return ""
    return table.rows[0][col].strip()


def column_names(table: Table, has_header: bool) -> list[str]:
    names = []
    for i in range(table.n_cols):
        h = header_cell(table, has_header, i)
        names.append(f"{column_letter(i)} — {h}" if h else f"Столбец {column_letter(i)}")
    return names


def with_names(table: Table, spec: TableSpec) -> TableSpec:
    return replace(
        spec,
        sheet=table.sheet,
        text_name=header_cell(table, spec.has_header, spec.text_col),
        title_name=header_cell(table, spec.has_header, spec.title_col) if spec.title_col is not None else "",
    )


def _is_number(value: str) -> bool:
    """Число, дата или время: только цифры и знаки вроде . , : - /."""
    return any(ch.isdigit() for ch in value) and all(ch.isdigit() or ch in " .,:-/+" for ch in value)


def guess_has_header(rows: list[list[str]]) -> bool:
    """Первая строка — заголовки, если на это есть явные признаки:
    над столбцом чисел стоит слово или она заметно короче данных в столбце.
    Длинная ячейка или ячейка с точкой в конце — это уже текст, а не заголовок."""
    if len(rows) < 2:
        return False
    first = [c.strip() for c in rows[0]]
    if not any(first) or any(len(c) > 60 or c[-1:] in ".!?…" for c in first if c):
        return False
    data = rows[1:51]
    for i, head in enumerate(first):
        values = [r[i].strip() for r in data if i < len(r) and r[i].strip()]
        if not head or not values:
            continue
        if not _is_number(head) and all(_is_number(v) for v in values):
            return True
        lengths = sorted(len(v) for v in values)
        if len(head) * 2 <= lengths[len(lengths) // 2]:
            return True
    return False


def guess_spec(table: Table, has_header: bool | None = None) -> TableSpec:
    if has_header is None:
        has_header = guess_has_header(table.rows)
    data = table.rows[1:] if has_header else table.rows
    sample = data[:300]
    n = table.n_cols
    # столбец с текстом — в среднем больше всего букв на строку
    scores = []
    for i in range(n):
        letters = sum(sum(ch.isalpha() for ch in r[i]) for r in sample if i < len(r))
        scores.append(letters / max(1, len(sample)))
    text_col = max(range(n), key=scores.__getitem__) if n else 0

    title_col = None
    if has_header and table.rows:
        headers = [re.findall(r"\w+|№", c.lower()) for c in table.rows[0]]
        for hint in TITLE_HINTS:
            title_col = next(
                (i for i, words in enumerate(headers) if i != text_col and any(w.startswith(hint) for w in words)),
                None,
            )
            if title_col is not None:
                break
    return with_names(table, TableSpec(text_col, title_col, has_header))


def reuse_spec(table: Table, template: TableSpec) -> TableSpec | None:
    """Тот же выбор для другой таблицы: столбцы ищутся по заголовкам (без заголовков — по номерам).
    None — если столбец с текстом не найден и выбор нужно сделать заново."""
    if not table.rows:
        return None
    header = [c.strip().lower() for c in table.rows[0]] if template.has_header else []

    def find(name: str, index: int | None) -> int | None:
        if index is None:
            return None
        if template.has_header and name:
            return header.index(name.lower()) if name.lower() in header else None
        return index if index < table.n_cols else None

    text_col = find(template.text_name, template.text_col)
    if text_col is None:
        return None
    title_col = find(template.title_name, template.title_col)
    return with_names(table, TableSpec(text_col, title_col, template.has_header))


def resolve_column(table: Table, has_header: bool, value: str) -> int:
    """Столбец по заголовку, букве (A, B, AA) или номеру с 1 — для командной строки."""
    v = value.strip()
    if has_header and table.rows:
        for i, cell in enumerate(table.rows[0]):
            if cell.strip().lower() == v.lower():
                return i
    if v.isdigit() and 1 <= int(v) <= table.n_cols:
        return int(v) - 1
    if v.isalpha() and v.isascii():
        for i in range(table.n_cols):
            if column_letter(i) == v.upper():
                return i
    raise ValueError(f"столбец «{value}» не найден; есть: {', '.join(column_names(table, has_header))}")


# --- тексты ------------------------------------------------------------------

def table_texts(table: Table, spec: TableSpec) -> list[tuple[int, str | None, str]]:
    """[(номер строки в файле, название или None, текст), …]; строки с пустым текстом пропускаются."""
    out = []
    start = 1 if spec.has_header else 0
    for idx in range(start, len(table.rows)):
        row = table.rows[idx]
        text = row[spec.text_col].strip() if spec.text_col < len(row) else ""
        if not text:
            continue
        title = None
        if spec.title_col is not None and spec.title_col < len(row):
            title = row[spec.title_col].strip() or None
        out.append((table.first_row + idx, title, text))
    return out


def describe_spec(table: Table, spec: TableSpec) -> str:
    names = column_names(table, spec.has_header)
    text = f"текст — {names[spec.text_col]}" if spec.text_col < len(names) else "текст — ?"
    if spec.title_col is not None and spec.title_col < len(names):
        return f"{text}, название — {names[spec.title_col]}"
    return f"{text}, название — номер строки"


def table_items(
    name: str, table: Table, texts: list[tuple[int, str | None, str]]
) -> list[tuple[str, str, str]]:
    """Тексты таблицы в виде (название, текст, источник) для анализа."""
    where = str(table.path)
    if table.sheet and len(table.sheet_names) > 1:
        where += f", лист «{table.sheet}»"

    def item_name(row: int, title: str | None) -> str:
        if not title:
            return f"{name} — строка {row}"
        # одиночный номер из столбца «ID»/«№» сам по себе ничего не говорит
        return f"{name} — {title}" if _is_number(title) else title

    return [(item_name(row, title), text, f"{where}, строка {row}") for row, title, text in texts]
