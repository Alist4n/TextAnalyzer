"""DOCX-отчёт (python-docx). Содержит те же разделы, что и HTML, но без интерактива и деревьев."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..config import NRC_HOME
from ..nlp.emotions import EMOTION_RU, EMOTIONS, NEUTRAL
from ..nlp.models import CorpusResult, TextResult
from ..nlp.syntax import format_bases
from . import charts
from .common import CorpusOverview, corpus_overview

MAX_SENTENCES = 300


def _pct(v: float | None, digits: int = 0) -> str:
    return f"{(v or 0) * 100:.{digits}f}%"


def _shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _table(doc, header: Sequence[str], rows: Iterable[Sequence], widths: Sequence[float] | None = None, font_size: float = 9.5):
    rows = list(rows)
    n_cols = len(header)
    table = doc.add_table(rows=len(rows) + 1, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    # список ячеек строим один раз: table.add_row() и row.cells для каждой строки
    # заново обходят всю таблицу, и на сотнях строк это занимает минуты
    cells = table._cells

    def put(cell, value, bold: bool = False) -> None:
        run = cell.paragraphs[0].add_run("" if value is None else str(value))
        run.font.size = Pt(font_size)
        run.bold = bold

    for c, title in enumerate(header):
        put(cells[c], title, bold=True)
        _shade(cells[c], "E8EEF7")
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(list(row)[:n_cols]):
            put(cells[r * n_cols + c], value)
    if widths:
        for i, cell in enumerate(cells):
            cell.width = Cm(widths[i % n_cols])
    doc.add_paragraph()
    return table


def _picture(doc, png: bytes | None, width_cm: float = 16.0) -> None:
    if png:
        doc.add_picture(io.BytesIO(png), width=Cm(width_cm))


def _heading(doc, text: str, level: int) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x26, 0x32, 0x3F)


def _stats_section(doc, t: TextResult) -> None:
    s = t.stats
    _table(
        doc,
        ["Показатель", "Значение"],
        [
            ("Слов", s.n_words),
            ("Предложений", s.n_sentences),
            ("Абзацев", s.n_paragraphs),
            ("Символов (без пробелов)", f"{s.n_chars} ({s.n_chars_no_spaces})"),
            ("Уникальных лемм", s.n_unique_lemmas),
            ("Средняя длина предложения, слов", s.avg_sentence_len),
            ("Средняя длина слова, букв", s.avg_word_len),
            ("Слогов в слове (сред.)", s.avg_syllables),
            ("Лексическое разнообразие TTR / MATTR", f"{s.ttr} / {s.mattr}"),
            ("Доля слов, встретившихся один раз", _pct(s.hapax_ratio, 1)),
            ("Лексическая плотность", _pct(s.lexical_density, 1)),
        ],
        widths=[9, 7],
    )


def _emotion_section(doc, t: TextResult) -> None:
    e = t.emotions
    if not e:
        return
    _heading(doc, "Эмоции", 2)
    lines = []
    if e.dominant:
        lines.append(f"Преобладающее состояние (модель): {EMOTION_RU[e.dominant]}.")
    if e.leading_emotion and e.profile:
        lines.append(f"Ведущая эмоция: {EMOTION_RU[e.leading_emotion]} ({_pct(e.profile[e.leading_emotion], 1)}).")
    if e.sentence_share:
        lines.append(f"Предложений с выраженной эмоцией: {_pct(1 - e.sentence_share.get(NEUTRAL, 0))}.")
    if e.lexicon_totals is not None:
        lex = EMOTION_RU[e.lexicon_dominant] if e.lexicon_dominant else "эмоциональная лексика не найдена"
        lines.append(f"По словарю преобладает: {lex}.")
    for line in lines:
        doc.add_paragraph(line)
    if e.profile:
        _picture(doc, charts.emotion_profile_chart(e.profile, "Эмоциональный профиль (модель)", "png"), 14)
        _picture(doc, charts.emotion_dynamics_chart(t.sentences, "png"), 16)
    if e.lexicon_totals and any(e.lexicon_totals.values()):
        _table(
            doc,
            ["Эмоция", "Сумма весов", "На 100 слов", "Частые слова"],
            [
                (
                    EMOTION_RU[emo],
                    e.lexicon_totals.get(emo, 0),
                    e.lexicon_density.get(emo, 0) if e.lexicon_density else "",
                    ", ".join(f"{w} ({c})" for w, c in (e.lexicon_top or {}).get(emo, [])[:10]),
                )
                for emo in EMOTIONS
            ],
            widths=[3, 2.5, 2.5, 8],
        )
    if e.lexicon_negated:
        doc.add_paragraph("С отрицанием (не учитывались): " + ", ".join(f"{w} ({c})" for w, c in e.lexicon_negated))


def _morphology_section(doc, t: TextResult, max_rows: int) -> None:
    if not t.pos_distribution:
        return
    _heading(doc, "Морфология", 2)
    _picture(doc, charts.bar_chart(t.pos_distribution, "Части речи", "png"), 14)
    total = sum(t.pos_distribution.values()) or 1
    _table(doc, ["Часть речи", "Кол-во", "Доля"],
           [(k, v, _pct(v / total, 1)) for k, v in t.pos_distribution.items()], widths=[8, 3, 3])
    doc.add_paragraph(f"Морфологическая омонимия: {_pct(t.ambiguity_rate, 1)} словоформ имеют несколько вариантов разбора.")

    if t.feature_distributions:
        _heading(doc, "Грамматические категории", 3)
        rows = [(title, ", ".join(f"{v}: {c}" for v, c in dist.items())) for title, dist in t.feature_distributions.items()]
        _table(doc, ["Категория", "Распределение"], rows, widths=[4.5, 11.5])

    if t.top_lemmas:
        _heading(doc, "Частотный словарь", 3)
        cols = [("NOUN", "Существительные"), ("VERB", "Глаголы"), ("ADJ", "Прилагательные"), ("ADV", "Наречия")]
        cols = [(k, title) for k, title in cols if k in t.top_lemmas]
        depth = max(len(t.top_lemmas[k]) for k, _ in cols) if cols else 0
        rows = []
        for i in range(depth):
            row = []
            for k, _ in cols:
                items = t.top_lemmas[k]
                row.append(f"{items[i][0]} ({items[i][1]})" if i < len(items) else "")
            rows.append(row)
        if cols:
            _table(doc, [title for _, title in cols], rows)

    if t.unknown_words:
        doc.add_paragraph("Слова вне словаря: " + ", ".join(w for w, _ in t.unknown_words))

    tokens = [tok for s in t.sentences for tok in s.tokens if not tok.is_punct]
    if tokens:
        _heading(doc, "Морфологический разбор по словам", 3)
        if len(tokens) > max_rows:
            doc.add_paragraph(f"Показаны первые {max_rows} слов из {len(tokens)}; полный разбор — в HTML/JSON-отчёте.")
        _table(
            doc,
            ["Словоформа", "Лемма", "Часть речи", "Признаки", "Член предл."],
            [(tok.text, tok.lemma, tok.pos_ru, tok.feats_ru, tok.member) for tok in tokens[:max_rows]],
            widths=[2.6, 2.6, 2.8, 6.0, 2.6],
            font_size=8,
        )


def _syntax_section(doc, t: TextResult) -> None:
    if not t.syntax_summary:
        return
    sx = t.syntax_summary
    _heading(doc, "Синтаксис", 2)
    _table(
        doc,
        ["Показатель", "Значение"],
        [
            ("Сложных предложений", _pct(sx["complex_share"])),
            ("Частей в предложении (сред.)", sx["avg_clauses"]),
            ("Глубина дерева: средняя / максимальная", f'{sx["avg_depth"]} / {sx["max_depth"]}'),
            ("Средняя длина зависимости", sx["mean_dep_distance"]),
            ("Причастных оборотов", sx["participial_count"]),
            ("Деепричастных оборотов", sx["adverbial_count"]),
            ("Рядов однородных членов", sx["homogeneous_count"]),
            ("Восклицательных предложений", _pct(sx["exclamatory_share"])),
        ]
        + [(f"Цель высказывания: {k}", v) for k, v in t.purpose_distribution.items()]
        + [(f"Строение: {k}", v) for k, v in t.structure_distribution.items()],
        widths=[10, 6],
    )


def _sentences_section(doc, t: TextResult) -> None:
    if not t.sentences:
        return
    _heading(doc, "Предложения", 2)
    sents = t.sentences[:MAX_SENTENCES]
    if len(t.sentences) > MAX_SENTENCES:
        doc.add_paragraph(f"Показаны первые {MAX_SENTENCES} предложений из {len(t.sentences)}.")
    rows = []
    for s in sents:
        kind = ", ".join(x for x in (s.purpose, "восклицательное" if s.exclamatory else "", s.structure) if x)
        bases = format_bases(s.bases)
        emo = ""
        if s.emotion_dominant:
            emo = f"{EMOTION_RU[s.emotion_dominant]} {_pct(s.emotion[s.emotion_dominant])}"
        lex = ", ".join(sorted({f"{h.token}→{EMOTION_RU[h.emotion]}" for h in s.lexicon_hits if not h.negated}))
        rows.append((s.index + 1, s.text, kind, bases, emo, lex))
    _table(doc, ["№", "Предложение", "Тип", "Грамм. основа", "Эмоция", "Эмоц. слова"], rows,
           widths=[0.9, 6.5, 2.8, 2.8, 1.8, 2.2], font_size=8)


def _entities_section(doc, t: TextResult, top_n: int) -> None:
    if not t.entities:
        return
    _heading(doc, "Именованные сущности", 2)
    _table(doc, ["Сущность", "Тип", "Упоминаний"],
           [(e.text, e.label_ru, e.count) for e in t.entities[: top_n * 3]], widths=[8, 4, 3])


def _overview_section(doc, ov: CorpusOverview) -> None:
    _heading(doc, "Все тексты вместе", 1)
    doc.add_paragraph(
        f"Текстов: {ov.n_texts}, слов: {ov.n_words}, предложений: {ov.n_sentences}. "
        f"В среднем {ov.avg_words} слов в тексте."
    )
    if ov.leading_counts:
        _picture(doc, charts.emotion_counts_chart(ov.leading_counts, f"Тексты по ведущей эмоции ({ov.leading_source})", "png"), 14)
    if ov.profile:
        _picture(doc, charts.emotion_profile_chart(ov.profile, "Средний эмоциональный профиль", "png"), 14)
    if ov.top_lemmas:
        _heading(doc, "Частые слова во всех текстах", 2)
        doc.add_paragraph(", ".join(f"{w} ({c})" for w, c in ov.top_lemmas))
    if ov.lexicon_top:
        _heading(doc, "Эмоциональные слова во всех текстах", 2)
        _table(
            doc,
            ["Эмоция", "Сумма весов", "Частые слова"],
            [
                (EMOTION_RU[e], (ov.lexicon_totals or {}).get(e, 0), ", ".join(f"{w} ({c})" for w, c in ov.lexicon_top.get(e, [])[:12]))
                for e in EMOTIONS
            ],
            widths=[3, 2.5, 10.5],
        )


def write_docx(corpus: CorpusResult, path: Path) -> Path:
    opts = corpus.options
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    doc.add_heading("Лингвистический анализ текста", level=0)
    doc.add_paragraph(
        f"Дата: {corpus.created.replace('T', ' ')}. Текстов: {len(corpus.texts)}. "
        f"Время анализа: {corpus.elapsed} с. TextAnalyzer {corpus.app_version}."
    )
    models = {k: v for k, v in corpus.models.items() if k not in ("lexicon_credit", "lexicon_note")}
    if models:
        doc.add_paragraph("Модели: " + "; ".join(f"{k}: {v}" for k, v in models.items()))
    if corpus.models.get("lexicon_credit"):
        doc.add_paragraph(corpus.models["lexicon_credit"] + " " + NRC_HOME)
    elif corpus.models.get("lexicon_note"):
        doc.add_paragraph(corpus.models["lexicon_note"])

    n = len(corpus.texts)
    detailed = max(0, int(opts.get("detailed_texts", 30)))
    if n > 1:
        _heading(doc, "Сводка по набору текстов", 1)
        rows = []
        for k, t in enumerate(corpus.texts, 1):
            e = t.emotions
            leading = "—"
            if e and e.leading_emotion:
                leading = f"{EMOTION_RU[e.leading_emotion]} ({_pct(e.profile[e.leading_emotion])})"
            elif e and e.profile:
                leading = "нейтрально"
            rows.append((
                k,
                t.name,
                t.stats.n_words,
                t.stats.n_sentences,
                t.stats.mattr,
                leading,
                EMOTION_RU[e.lexicon_dominant] if e and e.lexicon_dominant else "—",
            ))
        _table(doc, ["№", "Текст", "Слов", "Предл.", "MATTR", "Ведущая эмоция", "По словарю"],
               rows, font_size=8.5)
        if n <= detailed:
            _picture(doc, charts.corpus_emotions_chart(corpus.texts, "png"), 16)
        _overview_section(doc, corpus_overview(corpus, opts.get("top_n", 20)))
        if n > detailed:
            doc.add_paragraph(
                f"Подробные разделы построены для первых {detailed} текстов из {n} "
                "(настройка «Подробных разделов»); остальные тексты — в сводной таблице выше, "
                "а полный разбор каждого — в HTML- и JSON-отчётах."
            )

    for k, t in enumerate(corpus.texts, 1):
        if n > 1 and k > detailed:
            break
        if k > 1 or n > 1:
            doc.add_page_break()
        _heading(doc, f"{k}. {t.name}" if n > 1 else t.name, 1)
        if t.source:
            doc.add_paragraph(t.source).runs[0].font.size = Pt(9)
        _stats_section(doc, t)
        _emotion_section(doc, t)
        _morphology_section(doc, t, opts.get("max_token_rows_docx", 400))
        _syntax_section(doc, t)
        _sentences_section(doc, t)
        _entities_section(doc, t, opts.get("top_n", 20))

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
