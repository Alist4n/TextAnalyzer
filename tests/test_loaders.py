"""Какие файлы принимает программа и как они читаются."""

import pytest

from textanalyzer.loaders import UnsupportedFileError, collect_files, load_file, split_items, split_texts

TEXT = "Привет! Как дела? Я очень рад тебя видеть."


@pytest.mark.parametrize(
    "encoding, label",
    [
        ("utf-8", "UTF-8"),
        ("utf-8-sig", "UTF-8"),
        ("utf-16", "UTF-16"),
        ("cp1251", "Windows-1251"),
        ("koi8-r", "KOI8-R"),
        ("cp866", "CP866 (DOS)"),
    ],
)
def test_text_encodings_detected(tmp_path, encoding, label):
    path = tmp_path / "t.txt"
    path.write_bytes(TEXT.encode(encoding))
    loaded = load_file(path)
    assert loaded.text == TEXT
    assert loaded.encoding == label
    assert loaded.kind == "обычный текст"


@pytest.mark.parametrize(
    "encoding, label",
    [("cp1251", "Windows-1251"), ("koi8-r", "KOI8-R"), ("cp866", "CP866 (DOS)"), ("mac-cyrillic", "Mac Cyrillic")],
)
@pytest.mark.parametrize("text", ["Привет, мир.", "Да", "Ёлка — «ёж»! Съешь же ещё этих мягких французских булок."])
def test_short_russian_texts_in_single_byte_encodings(tmp_path, encoding, label, text):
    if encoding == "mac-cyrillic" and len(text) < 5:
        pytest.skip("«Да» в Mac Cyrillic побайтно совпадает с «„а» в Windows-1251 — различить нельзя")
    path = tmp_path / "t.txt"
    path.write_bytes(text.encode(encoding, errors="replace"))
    loaded = load_file(path)
    assert loaded.text == text.encode(encoding, errors="replace").decode(encoding)
    assert loaded.encoding == label


def test_markdown_is_read_as_plain_text(tmp_path):
    path = tmp_path / "t.md"
    path.write_text("# Заголовок\n\nТекст *с разметкой*.", encoding="utf-8")
    loaded = load_file(path)
    assert loaded.kind == "Markdown"
    assert "*с разметкой*" in loaded.text


def test_docx_paragraphs(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("Первый абзац.")
    document.add_paragraph("Второй абзац.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Текст в таблице"
    path = tmp_path / "t.docx"
    document.save(str(path))
    loaded = load_file(path)
    assert loaded.text.splitlines() == ["Первый абзац.", "Второй абзац."]
    assert loaded.encoding is None


def test_broken_docx_explained(tmp_path):
    path = tmp_path / "fake.docx"
    path.write_bytes(b"not a zip at all")
    with pytest.raises(UnsupportedFileError, match="не документ Word"):
        load_file(path)


@pytest.mark.parametrize("name, hint", [("a.pdf", "PDF"), ("a.doc", "Word 97"), ("a.xyz", ".xyz")])
def test_unsupported_formats_explained(tmp_path, name, hint):
    path = tmp_path / name
    path.write_bytes(b"data")
    with pytest.raises(UnsupportedFileError, match=hint):
        load_file(path)
    assert collect_files([path]).unsupported == [path]


# --- несколько текстов в одном файле ---------------------------------------------

def test_split_with_titles():
    text = "=== Отзыв 1\nХорошо.\n=== Отзыв 2 ===\nПлохо.\n=====\nБез названия."
    assert split_texts(text, "===") == [("Отзыв 1", "Хорошо."), ("Отзыв 2", "Плохо."), (None, "Без названия.")]


def test_split_keeps_preamble_and_skips_empty_parts():
    text = "Вступление.\n===\n\n===\nВторой.\n=== Пустой\n   \n"
    assert split_texts(text, "===") == [(None, "Вступление."), (None, "Второй.")]


def test_no_separator_means_one_text():
    text = "Обычный текст.\n\nВторой абзац."
    assert split_texts(text, "===") == [(None, text)]
    assert split_texts("=== А\nтекст", None) == [(None, "=== А\nтекст")]
    assert split_texts("=== А\nтекст", "  ") == [(None, "=== А\nтекст")]


def test_custom_separators():
    assert split_texts("*** Глава 1\nРаз.\n***\nДва.", "***") == [("Глава 1", "Раз."), (None, "Два.")]
    # буквы в названии не срезаются, даже если они есть в разделителе
    assert split_texts("Текст: тест\nРаз.", "Текст:") == [("тест", "Раз.")]


def test_separator_only_at_line_start():
    text = "Первый текст, в нём есть === посреди строки.\n=== Второй\nВторой текст."
    assert [t for t, _ in split_texts(text, "===")] == [None, "Второй"]


def test_split_items_names_and_sources():
    items = [
        ("отзывы", "=== Первый\nА.\n===\nБ.", "C:/отзывы.txt"),
        ("одиночный", "Просто текст.", "C:/одиночный.txt"),
        ("вставка", "===\nВ.\n===\nГ.", None),
    ]
    assert split_items(items, "===") == [
        ("Первый", "А.", "C:/отзывы.txt, текст 1 из 2"),
        ("отзывы — 2", "Б.", "C:/отзывы.txt, текст 2 из 2"),
        ("одиночный", "Просто текст.", "C:/одиночный.txt"),
        ("вставка — 1", "В.", "вставка, текст 1 из 2"),
        ("вставка — 2", "Г.", "вставка, текст 2 из 2"),
    ]
    assert split_items(items, None) == items


def test_folder_collection(tmp_path):
    (tmp_path / "a.txt").write_text("а", encoding="utf-8")
    (tmp_path / "b.docx").write_bytes(b"")
    (tmp_path / "~$b.docx").write_bytes(b"")      # временный файл Word — игнорируется
    (tmp_path / "c.pdf").write_bytes(b"")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "d.md").write_text("д", encoding="utf-8")
    selection = collect_files([tmp_path])
    assert [p.name for p in selection.files] == ["a.txt", "b.docx", "d.md"]
    assert selection.skipped_in_folders == 1
    assert selection.unsupported == []
