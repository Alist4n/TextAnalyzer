"""Таблицы CSV / Excel: чтение, угадывание столбцов, строки -> тексты."""

from datetime import datetime

import pytest

from textanalyzer.loaders import UnsupportedFileError, collect_files, is_table
from textanalyzer.tables import (
    TableSpec,
    column_letter,
    guess_spec,
    read_table,
    resolve_column,
    reuse_spec,
    table_items,
    table_texts,
    with_names,
)


def write_cp1251_csv(path):
    path.write_bytes(
        "ID;Товар;Отзыв;Оценка\n"
        "1;Кофемашина;Очень довольна, кофе отличный!;5\n"
        "2;Наушники;\"Сломались через неделю.\nПоддержка молчит; деньги не вернули.\";1\n"
        "3;Книга;;4\n".encode("cp1251")
    )
    return path


def test_excel_style_csv(tmp_path):
    table = read_table(write_cp1251_csv(tmp_path / "отзывы.csv"))
    assert table.encoding == "Windows-1251" and table.delimiter == ";"
    spec = guess_spec(table)
    assert spec.has_header and spec.text_col == 2 and spec.title_col == 1
    assert spec.text_name == "Отзыв" and spec.title_name == "Товар"
    texts = table_texts(table, spec)
    assert texts == [
        (2, "Кофемашина", "Очень довольна, кофе отличный!"),
        (3, "Наушники", "Сломались через неделю.\nПоддержка молчит; деньги не вернули."),
    ]  # строка 4 с пустым отзывом пропущена
    items = table_items("отзывы", table, texts)
    assert items[0][0] == "Кофемашина"
    assert items[1][2].endswith("отзывы.csv, строка 3")


def test_crlf_inside_quoted_cell(tmp_path):
    path = tmp_path / "t.csv"
    path.write_bytes("Товар;Отзыв\r\nНаушники;\"Сломались.\r\nПоддержка молчит.\"\r\n".encode("utf-8"))
    table = read_table(path)
    assert table_texts(table, guess_spec(table)) == [(2, "Наушники", "Сломались.\nПоддержка молчит.")]


def test_utf8_bom_comma_csv(tmp_path):
    path = tmp_path / "t.csv"
    path.write_bytes("﻿name,review\nАнна,\"Хорошо, но дорого\"\nБорис,\"Ужасно, никогда больше\"\n".encode("utf-8"))
    table = read_table(path)
    spec = guess_spec(table)
    assert table.delimiter == ","
    assert table_texts(table, spec) == [(2, "Анна", "Хорошо, но дорого"), (3, "Борис", "Ужасно, никогда больше")]


def test_single_column_without_header(tmp_path):
    # запятые внутри текста не должны превращаться в столбцы, а первая строка — в заголовок
    path = tmp_path / "t.csv"
    path.write_text("Первый длинный отзыв о товаре.\nВторой отзыв, тоже длинный.\n", encoding="utf-8")
    table = read_table(path)
    spec = guess_spec(table)
    assert table.delimiter is None and table.n_cols == 1
    assert not spec.has_header
    assert [t for _, _, t in table_texts(table, spec)] == ["Первый длинный отзыв о товаре.", "Второй отзыв, тоже длинный."]


def test_two_columns_without_header(tmp_path):
    path = tmp_path / "t.csv"
    path.write_text("Анна;Хорошо, но дорого\nБорис;Ужасно, никогда больше\nВера;Нормально\n", encoding="utf-8")
    table = read_table(path)
    spec = guess_spec(table)
    assert not spec.has_header and spec.text_col == 1
    assert len(table_texts(table, spec)) == 3


def test_single_column_with_header(tmp_path):
    path = tmp_path / "t.csv"
    path.write_text("Комментарий\nВсё отлично, спасибо большое за помощь!\nНе понравилось, долго ждал.\n", encoding="utf-8")
    table = read_table(path)
    spec = guess_spec(table)
    assert spec.has_header
    assert [row for row, _, _ in table_texts(table, spec)] == [2, 3]


def make_workbook(path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Пусто"
    ws = wb.create_sheet("Отзывы")
    ws.append([None])
    ws.append(["№", "Дата", "Комментарий респондента"])
    ws.append([1, datetime(2026, 10, 1), "Мне всё понравилось, спасибо!"])
    ws.append([2.0, datetime(2026, 10, 2, 14, 30), "Ничего не понял, очень запутанно."])
    wb.active = 1
    wb.save(path)
    return path


def test_xlsx_active_sheet_and_values(tmp_path):
    table = read_table(make_workbook(tmp_path / "анкета.xlsx"))
    assert table.sheet == "Отзывы" and table.sheet_names == ["Пусто", "Отзывы"]
    assert table.first_row == 2  # первая строка листа пустая
    assert table.rows[1][:2] == ["1", "2026-10-01"]
    assert table.rows[2][:2] == ["2", "2026-10-02 14:30"]
    spec = guess_spec(table)
    assert spec.text_col == 2 and spec.title_col == 0
    texts = table_texts(table, spec)
    assert [row for row, _, _ in texts] == [3, 4]  # номера строк как в Excel
    names = [name for name, _, _ in table_items("анкета", table, texts)]
    assert names == ["анкета — 1", "анкета — 2"]  # одиночный номер — с именем файла


def test_xlsx_other_sheet(tmp_path):
    table = read_table(make_workbook(tmp_path / "a.xlsx"), sheet="Пусто")
    assert table.sheet == "Пусто" and table.rows == []


def test_broken_xlsx(tmp_path):
    path = tmp_path / "fake.xlsx"
    path.write_bytes(b"not a zip")
    with pytest.raises(UnsupportedFileError, match="не книга Excel"):
        read_table(path)


def test_tables_are_collected_from_folders(tmp_path):
    (tmp_path / "a.csv").write_text("x", encoding="utf-8")
    (tmp_path / "b.xlsx").write_bytes(b"")
    (tmp_path / "~$b.xlsx").write_bytes(b"")
    (tmp_path / "c.xls").write_bytes(b"")
    selection = collect_files([tmp_path])
    assert [p.name for p in selection.files] == ["a.csv", "b.xlsx"]
    assert selection.skipped_in_folders == 1
    assert all(is_table(p) for p in selection.files)


def test_reuse_spec_by_header_names(tmp_path):
    first = read_table(write_cp1251_csv(tmp_path / "1.csv"))
    template = guess_spec(first)
    other = tmp_path / "2.csv"
    other.write_text("Отзыв;Оценка;Товар\nХороший чайник, рекомендую всем.;5;Чайник\n", encoding="utf-8")
    spec = reuse_spec(read_table(other), template)
    assert spec is not None and spec.text_col == 0 and spec.title_col == 2

    missing = tmp_path / "3.csv"
    missing.write_text("Комментарий;Товар\nНормально, без претензий к качеству.;Чайник\n", encoding="utf-8")
    assert reuse_spec(read_table(missing), template) is None  # столбца «Отзыв» нет — спросить заново


def test_resolve_column(tmp_path):
    table = read_table(write_cp1251_csv(tmp_path / "t.csv"))
    assert resolve_column(table, True, "отзыв") == 2
    assert resolve_column(table, True, "C") == 2
    assert resolve_column(table, True, "3") == 2
    with pytest.raises(ValueError, match="есть: A — ID"):
        resolve_column(table, True, "Комментарий")


def test_column_letters():
    assert [column_letter(i) for i in (0, 25, 26, 701, 702)] == ["A", "Z", "AA", "ZZ", "AAA"]


def test_spec_without_title(tmp_path):
    table = read_table(write_cp1251_csv(tmp_path / "t.csv"))
    spec = with_names(table, TableSpec(text_col=2, title_col=None, has_header=True))
    names = [n for n, _, _ in table_items("t", table, table_texts(table, spec))]
    assert names == ["t — строка 2", "t — строка 3"]
